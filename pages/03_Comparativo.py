"""
Página 03 — Comparativo por Red de Salud
Ranking de cobertura por Red para cada indicador con línea de meta.
"""
import io
import streamlit as st
import plotly.graph_objects as go
import pandas as pd
import numpy as np
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from utils.auth import require_auth
from utils.ui import load_css, render_sidebar_brand, render_sidebar_logout
from utils.constants import SEMAFORO, COLORS, INDICADORES

st.set_page_config(
    page_title='Comparativo por Red',
    page_icon='📊',
    layout='wide',
)
load_css()
require_auth()

fichas = st.session_state.get('fichas', {})

# ── Sidebar ───────────────────────────────────────────────────────────────────
with st.sidebar:
    render_sidebar_brand()
    if fichas:
        ids = sorted(fichas.keys())
        st.markdown('<p class="sb-section-title">🔍 INDICADOR</p>',
                    unsafe_allow_html=True)
        fid = st.selectbox(
            'Seleccionar indicador:',
            ids,
            key='comp_ficha',
            format_func=lambda x: f'{fichas[x]["icono"]} ID {x} — {fichas[x]["titulo"][:42]}',
        )
    render_sidebar_logout()

if not fichas:
    st.warning('⚠️ Primero carga los archivos Excel en la página de Inicio.')
    st.stop()

ficha   = fichas[fid]
df_base = ficha['df']
logro   = ficha.get('logro')
tipo    = ficha.get('tipo', 'pct')
unidad  = ficha.get('unidad', '%')

meta_escalonada       = INDICADORES.get(fid, {}).get('meta_escalonada', None)
tiene_meta_escalonada = meta_escalonada is not None


def get_meta_escalonada(den: int):
    for tier in meta_escalonada:
        if tier['den_min'] <= den <= tier['den_max']:
            return tier['umbral'], tier['logro']
    return meta_escalonada[-1]['umbral'], meta_escalonada[-1]['logro']


# ── Header ────────────────────────────────────────────────────────────────────
st.markdown(f"""
<div class="header-diresa">
  <div style="flex:1">
    <div style="font-size:0.85rem;color:#FFB703;font-weight:700;letter-spacing:0.08em;
                text-transform:uppercase;margin-bottom:4px;">
      Indicador {fid} &nbsp;·&nbsp; Comparativo por Red de Salud
    </div>
    <h1 style="font-size:1.55rem;font-weight:900;line-height:1.25;margin:0 0 10px 0;
               color:#ffffff;text-shadow:0 2px 8px rgba(0,0,0,0.5);">
      {ficha['icono']}&nbsp; {ficha['titulo']}
    </h1>
    <p style="font-size:0.82rem;opacity:0.65;margin:0;">
      DIRESA Huancavelica &nbsp;·&nbsp; 2026 &nbsp;·&nbsp; Solo MINSA
    </p>
  </div>
</div>""", unsafe_allow_html=True)

if 'red' not in df_base.columns or not df_base['red'].str.len().gt(0).any():
    if fid == '15':
        st.info(
            '🏥 **Ficha 15 — Mamografía bilateral de tamizaje**\n\n'
            'Este indicador se mide únicamente a nivel **departamental** '
            '(Hospital Regional de Huancavelica). No aplica comparativo por Red de Salud.'
        )
    elif fid == '16':
        st.info(
            '💉 **Ficha 16 — Vacuna VPH**\n\n'
            'La base de datos con desagregación por Red está pendiente de carga. '
            'Sube la nueva base cuando esté disponible.'
        )
    else:
        st.warning('⚠️ Este indicador no tiene datos de Red de Salud.')
    st.stop()

# ── Ficha 16: VPH con dos grupos de edad — mostrar dos gráficas ──────────────
sub_grupos = ficha.get('sub_grupos')
if sub_grupos:
    tab_labels = [sg['titulo'] for sg in sub_grupos]
    tabs_vph = st.tabs(tab_labels)
    for tab_vph, sg in zip(tabs_vph, sub_grupos):
        with tab_vph:
            df_g = df_base[df_base['categoria'] == sg['categoria']].copy()
            sg_thr = sg['logro'] * 100

            agg_g = (df_g[df_g['red'].str.len() > 0]
                     .groupby('red')
                     .agg(den=('den', 'sum'), num=('num', 'sum'))
                     .reset_index())
            agg_g['pct'] = np.where(agg_g['den'] > 0,
                                    agg_g['num'] / agg_g['den'] * 100, 0).round(1)

            den_g = int(df_g['den'].sum())
            num_g = int(df_g['num'].sum())
            pct_g = round(num_g / den_g * 100, 1) if den_g > 0 else 0

            # Colores ANTES del sort
            agg_g['color'] = agg_g['pct'].apply(
                lambda p: SEMAFORO['verde'] if p >= sg_thr
                else (SEMAFORO['amarillo'] if p >= sg_thr * 0.80 else SEMAFORO['rojo'])
            )
            agg_g = agg_g.sort_values('pct', ascending=False).reset_index(drop=True)
            agg_g['rank'] = range(1, len(agg_g) + 1)
            bc_g = agg_g['color'].tolist()

            BH = 0.275
            DX_G = 0.18
            ymx_raw = max(agg_g['pct'].max() if not agg_g.empty else 0, sg_thr, pct_g)
            DY_G = max(ymx_raw * 0.042, 2.0)
            ymx_g = max((ymx_raw + DY_G) * 1.30, 25)

            fg = go.Figure()
            fg.add_trace(go.Bar(
                name='Cobertura 2026',
                x=agg_g['red'], y=agg_g['pct'],
                marker=dict(color=bc_g, line=dict(color='rgba(255,255,255,0.25)', width=1.2)),
                text=[''] * len(agg_g), width=0.55,
                customdata=np.column_stack([agg_g['den'].values,
                                            agg_g['num'].values,
                                            agg_g['rank'].values]),
                hovertemplate=(
                    '<b>%{x}</b><br>Cobertura: <b>%{y:.1f}%</b><br>'
                    f'Meta: {sg_thr:.0f}%<br>'
                    'PROG: %{customdata[0]:,}<br>EJEC: %{customdata[1]:,}<br>'
                    'Ranking: #%{customdata[2]}<extra></extra>'
                ),
            ))

            for ix in range(len(agg_g)):
                hh = float(agg_g['pct'].iloc[ix])
                cc = bc_g[ix]
                if hh <= 0:
                    continue
                fg.add_shape(type='path',
                    path=(f'M {ix+BH},{0} L {ix+BH+DX_G},{DY_G} '
                          f'L {ix+BH+DX_G},{hh+DY_G} L {ix+BH},{hh} Z'),
                    xref='x', yref='y', fillcolor=_darken(cc, 0.48),
                    line=dict(color='rgba(0,0,0,0)', width=0), layer='above')
                fg.add_shape(type='path',
                    path=(f'M {ix-BH},{hh} L {ix+BH},{hh} '
                          f'L {ix+BH+DX_G},{hh+DY_G} L {ix-BH+DX_G},{hh+DY_G} Z'),
                    xref='x', yref='y', fillcolor=_lighten(cc, 1.35),
                    line=dict(color='rgba(0,0,0,0)', width=0), layer='above')

            fg.add_hline(y=pct_g, line_dash='dot', line_color='rgba(100,180,255,0.8)',
                         line_width=2,
                         annotation_text=f'  DIRESA: {pct_g:.1f}%',
                         annotation_position='top right',
                         annotation_font=dict(color='rgba(100,180,255,1)', size=12))
            fg.add_hline(y=sg_thr, line_dash='dash', line_color='#FFB703',
                         line_width=2.5,
                         annotation_text=f'  META: {sg_thr:.0f}%',
                         annotation_position='top left',
                         annotation_font=dict(color='#FFB703', size=14))

            for ix in range(len(agg_g)):
                hh = float(agg_g['pct'].iloc[ix])
                fg.add_annotation(
                    x=agg_g['red'].iloc[ix], y=hh + DY_G + ymx_g * 0.028,
                    xref='x', yref='y', text=f'<b>{hh:.1f}%</b>',
                    showarrow=False,
                    font=dict(size=15, color='white', family='Arial Black'),
                    bgcolor='rgba(0,0,0,0)', borderpad=0)

            fg.update_layout(
                plot_bgcolor='rgba(0,0,0,0)', paper_bgcolor='rgba(0,0,0,0)',
                font_color='white', height=450,
                margin=dict(l=10, r=10, t=20, b=80),
                xaxis=dict(title='Red de Salud',
                           gridcolor='rgba(255,255,255,0.05)',
                           tickfont=dict(size=11, color='white', family='Arial'),
                           tickangle=-15),
                yaxis=dict(title='% Cobertura', range=[0, ymx_g],
                           gridcolor='rgba(255,255,255,0.07)',
                           ticksuffix='%', tickfont=dict(size=11)),
                bargap=0.30, showlegend=False,
            )

            cc_d = (SEMAFORO['verde'] if pct_g >= sg_thr
                    else (SEMAFORO['amarillo'] if pct_g >= sg_thr * 0.80
                          else SEMAFORO['rojo']))
            ee_d = _emoji_from_color(cc_d)

            col_c, col_b = st.columns([5, 1])
            with col_b:
                st.markdown(f"""
<div style="background:#112240;border:2px solid #FFB703;border-radius:12px;
            padding:14px 10px;margin-top:28px;text-align:center;">
  <div style="color:#FFB703;font-weight:800;font-size:0.75rem;text-transform:uppercase;
              letter-spacing:0.06em;margin-bottom:4px;">📋 Resumen</div>
  <div style="color:#FFB703;font-size:0.82rem;margin:6px 0;">
    <b>META</b><br>
    <span style="font-size:1.6rem;font-weight:900;">{sg_thr:.0f}%</span>
  </div>
  <hr style="border-color:rgba(255,255,255,0.15);margin:8px 0;">
  <div style="color:rgba(255,255,255,0.8);font-size:0.78rem;margin:4px 0;">
    <b>LOGRO DIRESA</b><br>
    <span style="font-size:1.4rem;font-weight:900;color:{cc_d};">{pct_g:.1f}%</span> {ee_d}
  </div>
  <hr style="border-color:rgba(255,255,255,0.15);margin:8px 0;">
  <div style="color:rgba(255,255,255,0.5);font-size:0.72rem;line-height:1.7;">
    PROG<br><b style="color:white;">{den_g:,}</b><br>
    EJEC<br><b style="color:white;">{num_g:,}</b>
  </div>
</div>""", unsafe_allow_html=True)
            with col_c:
                st.plotly_chart(fg, use_container_width=True,
                                key=f'chart_vph_{sg["categoria"].replace(" ", "_")}')

            st.markdown(f'<div class="seccion-titulo">📋 Ranking — {sg["titulo"]}</div>',
                        unsafe_allow_html=True)
            tbl_g = agg_g[['rank', 'red', 'pct', 'den', 'num']].copy()
            tbl_g['pendiente'] = (tbl_g['den'] - tbl_g['num']).clip(lower=0)
            tbl_g['estado'] = tbl_g['pct'].apply(
                lambda p: '🟢 En meta' if p >= sg_thr
                else ('🟡 Cerca' if p >= sg_thr * 0.80 else '🔴 Bajo meta')
            )
            tbl_g.rename(columns={
                'rank': '#', 'red': 'Red de Salud', 'pct': 'Cobertura %',
                'den': 'PROG', 'num': 'EJEC', 'pendiente': 'Pendiente', 'estado': 'Estado',
            }, inplace=True)
            fila_d_g = {
                '#': '—', 'Red de Salud': '📊 DIRESA (Total)',
                'Cobertura %': pct_g, 'PROG': den_g, 'EJEC': num_g,
                'Pendiente': max(0, den_g - num_g),
                'Estado': ee_d + (' En meta' if cc_d == SEMAFORO['verde']
                          else (' Cerca' if cc_d == SEMAFORO['amarillo'] else ' Bajo meta')),
            }
            tbl_d_g = pd.concat([tbl_g, pd.DataFrame([fila_d_g])], ignore_index=True)
            st.dataframe(tbl_d_g, use_container_width=True, hide_index=True, column_config={
                '#':            st.column_config.TextColumn('#', width='small'),
                'Red de Salud': st.column_config.TextColumn('Red de Salud', width='large'),
                'Cobertura %':  st.column_config.NumberColumn('Cobertura %', format='%.1f%%', width='medium'),
                'PROG':         st.column_config.NumberColumn('PROG', format='%d', width='small'),
                'EJEC':         st.column_config.NumberColumn('EJEC', format='%d', width='small'),
                'Pendiente':    st.column_config.NumberColumn('Pendiente', format='%d', width='small'),
                'Estado':       st.column_config.TextColumn('Estado', width='medium'),
            })
    st.stop()

# ── Agrupación por RED ────────────────────────────────────────────────────────
df_con_red = df_base[df_base['red'].str.len() > 0]
agg = (df_con_red
       .groupby('red')
       .agg(den=('den', 'sum'), num=('num', 'sum'))
       .reset_index())
agg['pct'] = np.where(agg['den'] > 0, agg['num'] / agg['den'] * 100, 0).round(1)

den_total = int(df_base['den'].sum())
num_total = int(df_base['num'].sum())
pct_total = round(num_total / den_total * 100, 1) if den_total > 0 else 0
thr       = (logro or 0) * 100


# ── Helpers de color ──────────────────────────────────────────────────────────
def _color_fijo(pct_val: float) -> str:
    if tipo != 'pct' or logro is None:
        return '#4a85c0'
    if pct_val >= thr:
        return SEMAFORO['verde']
    elif pct_val >= thr * 0.80:
        return SEMAFORO['amarillo']
    else:
        return SEMAFORO['rojo']


def _color_escalonado(pct_val: float, den_val: int) -> str:
    umbral_r, logro_r = get_meta_escalonada(den_val)
    if pct_val >= logro_r * 100:
        return SEMAFORO['verde']
    elif pct_val >= umbral_r * 100:
        return SEMAFORO['amarillo']
    else:
        return SEMAFORO['rojo']


def _emoji_from_color(c: str) -> str:
    if c == SEMAFORO['verde']:    return '🟢'
    if c == SEMAFORO['amarillo']: return '🟡'
    if c == SEMAFORO['rojo']:     return '🔴'
    return '🔵'


def _hex_to_rgb(hx: str):
    hx = hx.lstrip('#')
    return tuple(int(hx[i:i+2], 16) for i in (0, 2, 4))


def _darken(hex_color: str, f: float = 0.48) -> str:
    r, g, b = _hex_to_rgb(hex_color)
    return f'rgb({int(r*f)},{int(g*f)},{int(b*f)})'


def _lighten(hex_color: str, f: float = 1.35) -> str:
    r, g, b = _hex_to_rgb(hex_color)
    return f'rgb({min(255,int(r*f))},{min(255,int(g*f))},{min(255,int(b*f))})'


# ── Calcular color por barra — ANTES del sort ─────────────────────────────────
if tiene_meta_escalonada:
    agg['color']      = [_color_escalonado(row['pct'], row['den'])
                         for _, row in agg.iterrows()]
    agg['logro_red']  = [get_meta_escalonada(d)[1] * 100 for d in agg['den']]
    agg['umbral_red'] = [get_meta_escalonada(d)[0] * 100 for d in agg['den']]
    agg['meta_txt']   = agg.apply(
        lambda r: f"Meta: {r['logro_red']:.0f}% (Umbral: {r['umbral_red']:.0f}%)", axis=1)
    u_diresa, l_diresa = get_meta_escalonada(den_total)
else:
    agg['color']      = [_color_fijo(p) for p in agg['pct']]
    agg['logro_red']  = thr
    agg['umbral_red'] = thr * 0.80 if thr > 0 else 0
    agg['meta_txt']   = f'Meta: {thr:.0f}%' if logro else 'Sin meta'

# Ordenar por cobertura descendente
agg = agg.sort_values('pct', ascending=False).reset_index(drop=True)
agg['rank'] = range(1, len(agg) + 1)
bar_colors  = agg['color'].tolist()

# ── Parámetros 3D ─────────────────────────────────────────────────────────────
BAR_HALF = 0.275   # mitad del ancho de barra (width=0.55)
DX       = 0.18    # profundidad horizontal 3D
y_max_raw = max(
    agg['pct'].max() if not agg.empty else 0,
    thr,
    pct_total,
    agg['logro_red'].max() if not agg.empty else 0,
)
DY   = max(y_max_raw * 0.042, 2.0)   # profundidad vertical 3D
y_max = (y_max_raw + DY) * 1.30
y_max = max(y_max, 25)

# ── Construcción del gráfico 3D ───────────────────────────────────────────────
fig = go.Figure()

# Barra principal (cara frontal)
fig.add_trace(go.Bar(
    name='Cobertura 2026',
    x=agg['red'],
    y=agg['pct'],
    marker=dict(
        color=bar_colors,
        line=dict(color='rgba(255,255,255,0.25)', width=1.2),
    ),
    text=[''] * len(agg),   # etiquetas vía annotations
    width=0.55,
    customdata=np.column_stack([
        agg['den'].values,
        agg['num'].values,
        agg['rank'].values,
        agg['logro_red'].values,
        agg['meta_txt'].values,
    ]),
    hovertemplate=(
        '<b>%{x}</b><br>'
        'Cobertura: <b>%{y:.1f}%</b><br>'
        '%{customdata[4]}<br>'
        'PROG: %{customdata[0]:,}<br>'
        'EJEC: %{customdata[1]:,}<br>'
        'Ranking: #%{customdata[2]}'
        '<extra></extra>'
    ),
))

# Caras 3D: cara derecha + cara superior por cada barra
for idx in range(len(agg)):
    h  = float(agg['pct'].iloc[idx])
    c  = bar_colors[idx]
    if h <= 0:
        continue

    # Cara derecha (sombra lateral — más oscura)
    fig.add_shape(
        type='path',
        path=(f'M {idx + BAR_HALF},{0} '
              f'L {idx + BAR_HALF + DX},{DY} '
              f'L {idx + BAR_HALF + DX},{h + DY} '
              f'L {idx + BAR_HALF},{h} Z'),
        xref='x', yref='y',
        fillcolor=_darken(c, 0.48),
        line=dict(color='rgba(0,0,0,0)', width=0),
        layer='above',
    )

    # Cara superior (brillo — más clara)
    fig.add_shape(
        type='path',
        path=(f'M {idx - BAR_HALF},{h} '
              f'L {idx + BAR_HALF},{h} '
              f'L {idx + BAR_HALF + DX},{h + DY} '
              f'L {idx - BAR_HALF + DX},{h + DY} Z'),
        xref='x', yref='y',
        fillcolor=_lighten(c, 1.35),
        line=dict(color='rgba(0,0,0,0)', width=0),
        layer='above',
    )

# Marcadores de meta escalonada
if tiene_meta_escalonada:
    for i, row in agg.iterrows():
        fig.add_trace(go.Scatter(
            x=[row['red']],
            y=[row['logro_red'] + DY + 1],
            mode='markers',
            marker=dict(symbol='line-ew', size=24, color='#FFB703',
                        line=dict(color='#FFB703', width=3)),
            name='Meta' if i == 0 else '',
            showlegend=(i == 0),
            hovertemplate=(
                f'<b>{row["red"]}</b><br>'
                f'Meta esperada: {row["logro_red"]:.0f}%<br>'
                f'Umbral mínimo: {row["umbral_red"]:.0f}%'
                '<extra></extra>'
            ),
        ))

# Línea DIRESA
fig.add_hline(
    y=pct_total,
    line_dash='dot',
    line_color='rgba(100,180,255,0.8)',
    line_width=2,
    annotation_text=f'  DIRESA: {pct_total:.1f}%',
    annotation_position='top right',
    annotation_font=dict(color='rgba(100,180,255,1)', size=12),
)

# Línea META fija
if logro and tipo == 'pct' and not tiene_meta_escalonada:
    fig.add_hline(
        y=thr,
        line_dash='dash',
        line_color='#FFB703',
        line_width=2.5,
        annotation_text=f'  META: {thr:.0f}%',
        annotation_position='top left',
        annotation_font=dict(color='#FFB703', size=14, family='Inter'),
    )

# Etiquetas de porcentaje grandes y visibles — encima de la cara superior
for idx in range(len(agg)):
    h = float(agg['pct'].iloc[idx])
    red_name = agg['red'].iloc[idx]
    fig.add_annotation(
        x=red_name,
        y=h + DY + y_max * 0.028,
        xref='x', yref='y',
        text=f'<b>{h:.1f}%</b>',
        showarrow=False,
        font=dict(size=15, color='white', family='Arial Black'),
        bgcolor='rgba(0,0,0,0)',
        borderpad=0,
    )

fig.update_layout(
    plot_bgcolor='rgba(0,0,0,0)',
    paper_bgcolor='rgba(0,0,0,0)',
    font_color='white',
    height=490,
    margin=dict(l=10, r=10, t=20, b=80),
    xaxis=dict(
        title='Red de Salud',
        gridcolor='rgba(255,255,255,0.05)',
        tickfont=dict(size=11, color='white', family='Arial'),
        tickangle=-15,
    ),
    yaxis=dict(
        title='% Cobertura',
        range=[0, y_max],
        gridcolor='rgba(255,255,255,0.07)',
        ticksuffix='%',
        tickfont=dict(size=11),
    ),
    bargap=0.30,
    showlegend=tiene_meta_escalonada,
    legend=dict(font=dict(color='white', size=10),
                bgcolor='rgba(0,0,0,0.3)', x=0.01, y=0.99),
)

# ── Caja resumen ──────────────────────────────────────────────────────────────
col_chart, col_box = st.columns([5, 1])

with col_box:
    if tiene_meta_escalonada:
        u_d, l_d     = get_meta_escalonada(den_total)
        color_diresa = _color_escalonado(pct_total, den_total)
        meta_box_html = f"""
  <div style="color:#FFB703;font-size:0.75rem;margin:4px 0;">
    <b>META (den={den_total:,})</b><br>
    <span style="font-size:1.3rem;font-weight:900;">{l_d*100:.0f}%</span>
  </div>
  <div style="color:rgba(255,255,255,0.5);font-size:0.7rem;margin:2px 0;">
    Umbral: {u_d*100:.0f}%
  </div>"""
        border_color = '#FFB703'
    elif logro and tipo == 'pct':
        color_diresa = _color_fijo(pct_total)
        meta_box_html = f"""
  <div style="color:#FFB703;font-size:0.82rem;margin:6px 0;">
    <b>META</b><br>
    <span style="font-size:1.6rem;font-weight:900;">{thr:.0f}%</span>
  </div>"""
        border_color = '#FFB703'
    else:
        color_diresa = '#4a85c0'
        meta_box_html = '<div style="color:rgba(255,255,255,0.4);font-size:0.75rem;">Sin meta fija</div>'
        border_color  = '#4a85c0'

    emoji_diresa = _emoji_from_color(color_diresa)
    logro_color  = ({'#2DC653': '#2DC653', '#FFB703': '#FFB703',
                     '#E63946': '#E63946'}.get(color_diresa, '#4a85c0'))

    st.markdown(f"""
<div style="background:#112240;border:2px solid {border_color};border-radius:12px;
            padding:14px 10px;margin-top:28px;text-align:center;">
  <div style="color:{border_color};font-weight:800;font-size:0.75rem;
              text-transform:uppercase;letter-spacing:0.06em;margin-bottom:4px;">
    📋 Resumen</div>
  {meta_box_html}
  <hr style="border-color:rgba(255,255,255,0.15);margin:8px 0;">
  <div style="color:rgba(255,255,255,0.8);font-size:0.78rem;margin:4px 0;">
    <b>LOGRO DIRESA</b><br>
    <span style="font-size:1.4rem;font-weight:900;color:{logro_color};">
      {pct_total:.1f}%
    </span> {emoji_diresa}
  </div>
  <hr style="border-color:rgba(255,255,255,0.15);margin:8px 0;">
  <div style="color:rgba(255,255,255,0.5);font-size:0.72rem;line-height:1.7;">
    PROG<br><b style="color:white;">{den_total:,}</b><br>
    EJEC<br><b style="color:white;">{num_total:,}</b>
  </div>
</div>""", unsafe_allow_html=True)

with col_chart:
    st.plotly_chart(fig, use_container_width=True, key='chart_comparativo')

# ── Nota meta escalonada ──────────────────────────────────────────────────────
if tiene_meta_escalonada:
    st.info("""
**📋 Meta escalonada — Depresión (Ficha 19)**

| Pacientes (PROG) | Umbral mínimo | **Logro esperado** |
|---|---|---|
| > 150 | 20% | **30%** |
| 101 – 150 | 30% | **40%** |
| 60 – 100 | 35% | **50%** |
| < 60 | 40% | **60%** |
""")

# ── Tabla de ranking ──────────────────────────────────────────────────────────
st.markdown('<div class="seccion-titulo">📋 Ranking de Redes — Tabla Resumen</div>',
            unsafe_allow_html=True)

tbl = agg[['rank', 'red', 'pct', 'den', 'num']].copy()
tbl['pendiente'] = (tbl['den'] - tbl['num']).clip(lower=0)

if tiene_meta_escalonada:
    tbl['meta_%'] = agg['logro_red'].apply(lambda x: f'{x:.0f}%')
    tbl['estado'] = [
        _emoji_from_color(c) + (
            ' En meta' if c == SEMAFORO['verde']
            else (' Cerca' if c == SEMAFORO['amarillo'] else ' Bajo meta')
        )
        for c in bar_colors
    ]
elif logro and tipo == 'pct':
    tbl['estado'] = tbl['pct'].apply(
        lambda p: '🟢 En meta' if p >= thr
        else ('🟡 Cerca' if p >= thr * 0.80 else '🔴 Bajo meta')
    )

tbl = tbl.rename(columns={
    'rank': '#', 'red': 'Red de Salud',
    'pct': 'Cobertura %', 'den': 'PROG', 'num': 'EJEC',
    'pendiente': 'Pendiente', 'meta_%': 'Meta', 'estado': 'Estado',
})

fila_d: dict = {
    '#': '—', 'Red de Salud': '📊 DIRESA (Total)',
    'Cobertura %': pct_total,
    'PROG': den_total, 'EJEC': num_total,
    'Pendiente': max(0, den_total - num_total),
}
if tiene_meta_escalonada:
    fila_d['Meta']   = f'{l_d*100:.0f}%'
    fila_d['Estado'] = emoji_diresa + (
        ' En meta' if color_diresa == SEMAFORO['verde']
        else (' Cerca' if color_diresa == SEMAFORO['amarillo'] else ' Bajo meta')
    )
elif logro and tipo == 'pct':
    c_d = _color_fijo(pct_total)
    fila_d['Estado'] = _emoji_from_color(c_d) + (
        ' En meta' if c_d == SEMAFORO['verde']
        else (' Cerca' if c_d == SEMAFORO['amarillo'] else ' Bajo meta')
    )

tbl_display = pd.concat([tbl, pd.DataFrame([fila_d])], ignore_index=True)

col_cfg = {
    '#':            st.column_config.TextColumn('#', width='small'),
    'Red de Salud': st.column_config.TextColumn('Red de Salud', width='large'),
    'Cobertura %':  st.column_config.NumberColumn('Cobertura %', format='%.1f%%', width='medium'),
    'PROG':         st.column_config.NumberColumn('PROG', format='%d', width='small'),
    'EJEC':         st.column_config.NumberColumn('EJEC', format='%d', width='small'),
    'Pendiente':    st.column_config.NumberColumn('Pendiente', format='%d', width='small'),
}
if 'Meta'   in tbl_display.columns: col_cfg['Meta']   = st.column_config.TextColumn('Meta',   width='small')
if 'Estado' in tbl_display.columns: col_cfg['Estado'] = st.column_config.TextColumn('Estado', width='medium')

st.dataframe(tbl_display, use_container_width=True, hide_index=True, column_config=col_cfg)

# ── Leyenda ───────────────────────────────────────────────────────────────────
if logro and tipo == 'pct' and not tiene_meta_escalonada:
    st.markdown(f"""
<div style="display:flex;gap:16px;flex-wrap:wrap;margin-top:8px;font-size:0.8rem;color:#8892a4;">
  <span>🟢 En meta (≥ {thr:.0f}%)</span>
  <span>🟡 Cerca (≥ {thr*0.8:.0f}%)</span>
  <span>🔴 Bajo meta (< {thr*0.8:.0f}%)</span>
  <span style="color:rgba(100,180,255,0.8);">― ― DIRESA total</span>
  <span style="color:#FFB703;">--- META {thr:.0f}%</span>
</div>""", unsafe_allow_html=True)
elif tiene_meta_escalonada:
    st.markdown("""
<div style="display:flex;gap:16px;flex-wrap:wrap;margin-top:8px;font-size:0.8rem;color:#8892a4;">
  <span>🟢 Alcanzó logro esperado</span>
  <span>🟡 Entre umbral y logro</span>
  <span>🔴 Por debajo del umbral</span>
  <span style="color:#FFB703;">— Meta individual por red</span>
  <span style="color:rgba(100,180,255,0.8);">― ― DIRESA total</span>
</div>""", unsafe_allow_html=True)

# ── Exportar ──────────────────────────────────────────────────────────────────
st.markdown('<div class="seccion-titulo">⬇️ Exportar</div>', unsafe_allow_html=True)

exp_col1, exp_col2 = st.columns(2)

with exp_col1:
    try:
        img_bytes = fig.to_image(format='png', width=1400, height=700, scale=2)
        st.download_button(
            label='🖼️ Descargar gráfico (PNG)',
            data=img_bytes,
            file_name=f'comparativo_red_{fid}_2026.png',
            mime='image/png',
            use_container_width=True,
        )
    except Exception:
        html_bytes = fig.to_html(full_html=True, include_plotlyjs='cdn').encode('utf-8')
        st.download_button(
            label='🖼️ Descargar gráfico (HTML)',
            data=html_bytes,
            file_name=f'comparativo_red_{fid}_2026.html',
            mime='text/html',
            use_container_width=True,
        )

with exp_col2:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        tbl_excel = tbl_display.copy()
        if 'Cobertura %' in tbl_excel.columns:
            tbl_excel['Cobertura %'] = pd.to_numeric(tbl_excel['Cobertura %'], errors='coerce')
        tbl_excel.to_excel(writer, index=False, sheet_name='Comparativo por Red', startrow=3)
        ws = writer.sheets['Comparativo por Red']
        ws['A1'] = f'DIRESA HUANCAVELICA — Comparativo por Red — Indicador {fid}'
        ws['A2'] = ficha['titulo']
        ws['A3'] = f'Año: 2026    PROG: {den_total:,}    EJEC: {num_total:,}    Cobertura DIRESA: {pct_total:.1f}%'
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 12
        ws['A1'].font = Font(bold=True, size=13)
        ws['A2'].font = Font(italic=True, size=10)
        ws['A3'].font = Font(size=10, color='444444')
        # Header row (row 4)
        hfill = PatternFill('solid', start_color='003087', end_color='003087')
        hfont = Font(bold=True, color='FFFFFF', size=11)
        for cell in ws[4]:
            cell.fill = hfill
            cell.font = hfont
            cell.alignment = Alignment(horizontal='center', vertical='center')
        # Data rows — bold text, alternating fill, centered numbers
        n_rows = len(tbl_excel)
        thin_side = Side(style='thin', color='CCCCCC')
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        fill_even = PatternFill('solid', start_color='EEF2FF', end_color='EEF2FF')
        fill_odd  = PatternFill('solid', start_color='FFFFFF', end_color='FFFFFF')
        for row_i in range(n_rows):
            excel_row = row_i + 5   # data starts at row 5 (1-indexed)
            row_fill = fill_even if row_i % 2 == 0 else fill_odd
            for col_i, cell in enumerate(ws[excel_row]):
                cell.font = Font(bold=True, size=10)
                cell.fill = row_fill
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal='center' if col_i != 1 else 'left',
                    vertical='center'
                )
        # DIRESA total summary row
        total_row = n_rows + 5
        total_data = ['', 'DIRESA HUANCAVELICA', den_total, num_total, f'{pct_total:.1f}%', '']
        for col_i, val in enumerate(total_data, start=1):
            c = ws.cell(row=total_row, column=col_i, value=val)
            c.font = Font(bold=True, size=11, color='FFFFFF')
            c.fill = PatternFill('solid', start_color='003087', end_color='003087')
            c.alignment = Alignment(horizontal='center' if col_i != 2 else 'left', vertical='center')
            c.border = thin_border
    st.download_button(
        label='📊 Descargar tabla (Excel)',
        data=output.getvalue(),
        file_name=f'comparativo_red_{fid}_2026.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        use_container_width=True,
    )
